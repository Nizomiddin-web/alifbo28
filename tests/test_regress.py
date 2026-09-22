"""Auditda topilgan va tuzatilgan nuqsonlar uçun regressiya testlari.

Har bir test nomidagi raqam alifbo auditidagi topilma raqami.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from yozuv import to_cyrillic, to_new, to_old  # noqa: E402
from yozuv.alphabets import OKINA  # noqa: E402

fails = 0
TU = "’"


def eq(got, want, label):
    global fails
    ok = got == want
    if not ok:
        fails += 1
    print(f"[{'OK ' if ok else 'XATO'}] {label}"
          + ("" if ok else f"\n        kutildi: {want!r}\n        natija : {got!r}"))


# ---- #1 / #9: /alifbo izohi render() da buzilmasin -----------------------
from bot.texts import render, t  # noqa: E402

for ui in ("yangi", "eski", "kirill"):
    note = t("alphabet_note", ui, size=28, old=f"O{OKINA}, G{OKINA}, Sh, Ch",
             new="Ö, Ğ, Ş, Ç", ng="ng", ng_misol="tong, köngil")
    eq(f"O{OKINA}, G{OKINA}, Sh, Ch" in note and "Ö, Ğ, Ş, Ç" in note and "«ng»" in note,
       True, f"#1 alifbo izohidagi harflar ögirilmaydi ({ui})")

# ---- #11 / #15: «eski» interfeys rejimi kanonik belgi berişi kerak -------
eq("o" + OKINA + "tkaz" in render("Men o'tkazaman", "eski"), True,
   "#11 «eski» rejimda ASCII apostrof emas, ʻ (U+02BB)")
eq("'" in render("Men o'tkazaman va so'zlayman", "eski"), False,
   "#15 «eski» rejimda ASCII apostrof qolmaydi")

# ---- #2 / #10: matn oxiridagi kirill «ъ» yöqolmasin ----------------------
eq(to_new("мавзуъ"), "mavzu" + TU, "#2 matn oxiridagi ъ saqlandi")
eq(to_new("тобеъ"), "tobe" + TU, "#2 тобеъ")
eq(to_new("манбаъ вазъ"), f"manba{TU} vaz{TU}", "#2 bir neça söz")
eq(to_new("объект"), "obyekt", "#2 ъ+е hamon tuşib qoladi")
eq(to_cyrillic(to_new("мавзуъ")), "мавзуъ", "#10 aylanma saqlandi")

# ---- #3 / #4: kirilldagi alohida harflar digrafga qöşilmasin -------------
eq(to_new("тасҳиҳ"), "tashih", "#3 сҳ yangi alifboda s+h")
eq(to_new("мусҳаф"), "mushaf", "#3 мусҳаф")
eq(to_old("тасҳиҳ"), f"tas{TU}hih", "#3 amaldagi alifboda tutuq bilan ajratiladi")
eq(to_old("Исҳоқ"), f"Is{TU}hoq", "#3 Исҳоқ → Is’hoq")
eq(to_new("гъ"), "g" + TU, "#4 гъ → g’ (ğ emas)")
eq(to_new("оъ"), "o" + TU, "#4 оъ → o’ (ö emas)")
eq(to_new("ёъ"), "yo" + TU, "#4 ёъ → yo’")
eq(to_new("Тошкент шаҳри"), "Toşkent şahri", "#4 oddiy ш hamon ş")

# ---- #5: lotin→kirill, unlidan keyingi «e» → «э» -------------------------
eq(to_cyrillic("aeroport"), "аэропорт", "#5 aeroport")
eq(to_cyrillic("poeziya"), "поэзия", "#5 poeziya")
eq(to_cyrillic("duet"), "дуэт", "#5 duet")
eq(to_cyrillic("keldi ketsa eshik"), "келди кетса эшик", "#5 undoshdan keyin ösmadi")
eq(to_cyrillic("Yevropa teatr ideal"), "Европа театр идеал", "#5 ye/undosh holatlari")
eq(to_new(to_cyrillic("aeroport")), "aeroport", "#5 aylanma barqaror")

# ---- #6: istisnolar kirill yönalişida lotinça qolmasin -------------------
eq(to_cyrillic("Ishoq va Ashob keldi"), "Исҳоқ ва Асҳоб келди", "#6 oddiy istisnolar ötkazildi")
eq(to_cyrillic("Chrome brauzeri"), "Chrome браузери", "#6 brend lotinça qoladi")

# ---- #7: istisno söz qöşimça bilan kelganda ------------------------------
for w in ("Chelseaning", "Chromeda", "Ishoqning", "Ashobdan", "Porscheda"):
    eq(to_new(w), w, f"#7 qöşimçali istisno: {w}")
eq(to_new("chelsi"), "çelsi", "#7 istisno bölmagan söz ötkaziladi")

# ---- #12: butunlay bosh harfli sözda tutuq → «Ъ» -------------------------
eq(to_cyrillic("SAN" + TU + "AT"), "САНЪАТ", "#12 SAN’AT → САНЪАТ")
eq(to_cyrillic("ma" + TU + "no"), "маъно", "#12 kiçik harfda ъ")
eq(to_new("САНЪАТ"), "SAN" + TU + "AT", "#12 teskari yönaliş")

# ---- #13: tirnoq içidagi söz oxiridagi «o’» yutilmasin -------------------
eq(to_new("'radio' va 'tango'"), "'radio' va 'tango'", "#13 ASCII tirnoq saqlandi")
eq(to_new("‘radio’"), "‘radio’", "#13 tipografik tirnoq saqlandi")
eq(to_new("tog' bog' sog'lom"), "toğ boğ soğlom", "#13 haqiqiy oʻ/gʻ ösgarmadi")
eq(to_new("tog', tog'. tog'!"), "toğ, toğ. toğ!", "#13 tinish belgisidan oldin")
eq(to_new("'so'z' deb"), "'söz' deb", "#13 tirnoq içidagi haqiqiy oʻ")
eq(to_new("tog" + OKINA), "toğ", "#13 aniq ʻ (U+02BB) har doim ğ/ö")

# ---- #14: yopuvçi tirnoq kirillda «ъ» bölmasin ---------------------------
eq(to_cyrillic("‘salom’ dedi"), "‘салом’ деди", "#14 yopuvçi tirnoq saqlandi")
eq(to_cyrillic("mavzu" + TU), "мавзуъ", "#14 haqiqiy tutuq hamon ъ")

# ==========================================================================
# Ikkinçi töp: auditning past ahamiyatli, ammo haqiqiy topilmalari
# ==========================================================================

# ---- eskirgan U+02BC hеç qayerda qolmasin --------------------------------
from yozuv import Script, detect  # noqa: E402

eq("\u02bc" in Script.UNKNOWN.value, False, "#L1 Script.UNKNOWN da eskirgan ʼ yöq")
eq(detect("12345"), Script.UNKNOWN, "#L1 detect hamon işlaydi")
eq(to_new("mavzu\u02bc"), "mavzu" + TU, "#L2 söz oxiridagi U+02BC kanonik ’ ga keladi")
eq(to_new("mavzu" + OKINA), "mavzu" + TU, "#L2 söz oxiridagi U+02BB ham")

# ---- ts → ц işonçli naqşlarda --------------------------------------------
for src, want in (("konstitutsiya", "конституция"), ("aktsiya", "акция"),
                  ("funktsiya", "функция"), ("revolyutsiya", "революция"),
                  ("natsional", "национал"), ("revolyutsion", "революцион")):
    eq(to_cyrillic(src), want, f"#L3 ts→ц: {src}")
for src, want in (("ketsa", "кетса"), ("otsa", "отса"), ("kutsa", "кутса"),
                  ("aytsa", "айтса"), ("yotsa", "ётса")):
    eq(to_cyrillic(src), want, f"#L3 yolğon ijobiy emas: {src}")

# ---- qöş «цц» -------------------------------------------------------------
eq(to_new("пицца"), "pitsa", "#L4 пицца → pitsa")
eq(to_cyrillic("pitssa"), "пицца", "#L4 teskari yönaliş")

# ---- undoşdan keyingi «ye» → «ъе» ----------------------------------------
eq(to_cyrillic("obyekt"), "объект", "#L5 obyekt → объект")
eq(to_cyrillic("subyekt"), "субъект", "#L5 subyekt")
eq(to_new(to_cyrillic("obyekt")), "obyekt", "#L5 aylanma barqaror")
eq(to_cyrillic("yer yetti"), "ер етти", "#L5 söz boşidagi ye ösmadi")
eq(to_cyrillic("dunyo quyosh"), "дунё қуёш", "#L5 yo ga tegmadi")

# ---- interfeys matnlari kirillda buzilmasin ------------------------------
eq(".txt" in t("file_bad", "kirill"), True, "#L6 fayl kengaytmalari saqlandi")
eq("README" in t("ocr_off", "kirill"), True, "#L6 README saqlandi")
eq("каналлар" in t("admin_stats", "kirill", users=0, today=0, week=0, subs=0, blocked=0,
                     conv=0, chars=0, chats=0),
   True, "#L6 «Guruh va kanallar» töliq ögirildi")

# ==========================================================================
# Uçinçi töp: konvertor qayta yozilgandan keyingi audit
# ==========================================================================

# ---- #A1 GUARD _caps_context ni kör qilib qöymasin -----------------------
eq(to_old("ТОШКЕНТ ШАҲРИ"), "TOSHKENT SHAHRI", "#A1 ALL-CAPS kirill → amaldagi alifbo")
eq(to_old("АҚШ"), "AQSH", "#A1 qisqartma")
eq(to_old("ИЧКИ ИШЛАР"), "ICHKI ISHLAR", "#A1 Ç va Ş birga")
eq(to_old("ЎЗБЕКИСТОН ЧЕМПИОНИ"), f"O{OKINA}ZBEKISTON CHEMPIONI", "#A1 Ö va Ç aralaş")
eq(to_old("Тошкент шаҳри"), "Toshkent shahri", "#A1 oddiy söz ösmadi")

# ---- #A2 bosh harfli söz ALL-CAPS deb hisoblanmasin ----------------------
eq(to_cyrillic("A" + TU + "zam va E" + TU + "lon"), "Аъзам ва Эълон", "#A2 bosh harf + tutuq")
eq(to_cyrillic("SAN" + TU + "AT"), "САНЪАТ", "#A2 haqiqiy ALL-CAPS hamon Ъ")
eq(to_new(to_cyrillic("A" + TU + "zam")), "A" + TU + "zam", "#A2 aylanma barqaror")

# ---- #A3 tirnoq juftligi tekşirilsin -------------------------------------
eq(to_new("«tog' va bog'»"), "«toğ va boğ»", "#A3 «» içidagi gʻ ögiriladi")
eq(to_cyrillic("«tog'»"), "«тоғ»", "#A3 «» kirillda")
eq(to_cyrillic("«mavzu" + TU + "»"), "«мавзуъ»", "#A3 «» içidagi tutuq")
eq(to_new('"tog\'"'), '"toğ"', "#A3 qöştirnoq apostrofni yopmaydi")
eq(to_new("\u2039tog'\u203a"), "\u2039toğ\u203a", "#A3 ‹› juftligi")
eq(to_new("\u201etog'\u201c"), "\u201etoğ\u201c", "#A3 „“ juftligi")
eq(to_new("'radio' va 'tango'"), "'radio' va 'tango'", "#A3 juft ASCII tirnoq saqlandi")
eq(to_new("\u2018radio\u2019"), "\u2018radio\u2019", "#A3 juft ‘’ saqlandi")

# ---- içki ajratgiç foydalanuvçi matnidan kelsa ---------------------------
from yozuv.converter import GUARD  # noqa: E402

dirty = f"tas{GUARD}hih va Tos{GUARD}hkent"
eq(to_new(dirty), "taşih va Toşkent", "#L7 GUARD kirişda natijaga taʼsir qilmaydi (yangi)")
eq(to_old(dirty), "tashih va Toshkent", "#L7 GUARD kirişda (amaldagi)")
eq(GUARD in to_new(dirty) or GUARD in to_old(dirty) or GUARD in to_cyrillic(dirty),
   False, "#L7 GUARD çiqişga sizib ötmaydi")

# ---- unumdorlik: 100 KB matn maqbul vaqtda ötadi -------------------------
import time  # noqa: E402

big = "Ўзбекистон Республикаси Тошкент шаҳри бўйича маълумот. " * 2000
_t0 = time.perf_counter()
to_new(big)
_dt = time.perf_counter() - _t0
eq(_dt < 3.0, True, f"#L8 100 KB kirill matn {_dt * 1000:.0f} ms da ötdi (< 3000 ms)")

# ---- #8: .docx run çegarasi ----------------------------------------------
try:
    import docx
    from bot.db import User
    from bot.handlers.documents import _convert_docx

    d = docx.Document()
    par = d.add_paragraph()
    for chunk in ["Tos", "hkent ", "c", "hiroq ", "o", OKINA + "quvchi ", "Chel", "sea"]:
        par.add_run(chunk)
    buf = io.BytesIO(); d.save(buf)
    out = docx.Document(io.BytesIO(_convert_docx(buf.getvalue(), User(user_id=1))))
    eq(out.paragraphs[0].text, "Toşkent çiroq öquvçi Chelsea", "#8 run çegarasidagi digraflar")

    d2 = docx.Document(); p2 = d2.add_paragraph()
    p2.add_run("Toshkent "); bold = p2.add_run("shahri"); bold.bold = True
    b2 = io.BytesIO(); d2.save(b2)
    o2 = docx.Document(io.BytesIO(_convert_docx(b2.getvalue(), User(user_id=1))))
    eq([(r.text, bool(r.bold)) for r in o2.paragraphs[0].runs],
       [("Toşkent ", False), ("şahri", True)], "#8 formatlaş saqlandi")
except ModuleNotFoundError:
    print("[SKIP] #8 — python-docx örnatilmagan")

print("\n" + ("BARÇASI OʻTDI" if not fails else f"{fails} ta test yiqildi"))
raise SystemExit(1 if fails else 0)
