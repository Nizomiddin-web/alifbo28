from __future__ import annotations

import io
import re
from pathlib import Path

from aiogram import F, Router
from aiogram.types import BufferedInputFile, Message

from .. import db
from ..config import config
from ..db import User
from ..texts import t
from ..utils import convert_for

router = Router(name="documents")

TEXT_EXT = {".txt", ".md", ".csv", ".srt", ".vtt", ".sub", ".json", ".ini", ".log"}
DOCX_EXT = {".docx"}


_PLAUSIBLE = re.compile(r"[0-9A-Za-z\u00c0-\u024f\u02bb\u02bc\u0400-\u04ff\s.,;:!?()\[\]{}\-–—«»\"'’/\\@#%&*+=_|<>№$€]")
_CANDIDATES = ("utf-8-sig", "cp1251", "cp1252", "koi8-r", "iso8859-5", "cp866", "latin-1")


def _score(text: str) -> float:
    """Matnning «oqilona» belgilar ulusi — kodlaşni tanlaş uçun."""
    if not text:
        return 0.0
    good = len(_PLAUSIBLE.findall(text))
    bad = text.count("\ufffd") * 5
    return (good - bad) / len(text)


def _decode(data: bytes) -> str:
    """Fayl baytlarini matnga aylantiradi (kodlaşni taxmin qilib)."""
    if not data:
        return ""
    try:
        return data.decode("utf-8-sig")          # eng keng tarqalgan hol
    except UnicodeDecodeError:
        pass

    candidates: list[tuple[float, int, str]] = []
    for rank, enc in enumerate(_CANDIDATES):
        try:
            text = data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
        candidates.append((_score(text), -rank, text))

    try:
        from charset_normalizer import from_bytes
        best = from_bytes(data).best()
        if best is not None:
            text = str(best)
            candidates.append((_score(text), -len(_CANDIDATES), text))
    except ModuleNotFoundError:
        pass

    if not candidates:
        return data.decode("utf-8", errors="replace")
    return max(candidates)[2]


def _convert_paragraph(paragraph, user: User) -> None:
    """Paragrafni BUTUN holda ötkazadi.

    Word bitta sözni bir neça `run` ga bölib taşlaşi mumkin ("Tos" + "hkent").
    Har bir run ni alohida ötkazsak, run çegarasidagi sh/ch/o` digraflari
    yöqoladi. Şu sababli avval butun paragrafni ötkazamiz; agar run lar
    böyiça ötkaziş aynan şu natijani bersa — formatlaşni saqlab, har bir run
    ni alohida yozamiz, aks holda butun matnni birinçi run ga joylaymiz.
    """
    runs = paragraph.runs
    if not runs:
        return
    whole = "".join(r.text for r in runs)
    if not whole.strip():
        return
    target = convert_for(user, whole).text
    per_run = [convert_for(user, r.text).text for r in runs]
    if "".join(per_run) == target:
        for run, text in zip(runs, per_run):
            if run.text != text:
                run.text = text
        return
    runs[0].text = target
    for run in runs[1:]:
        run.text = ""


def _convert_docx(data: bytes, user: User) -> bytes:
    import docx  # python-docx

    doc = docx.Document(io.BytesIO(data))

    def fix_paragraphs(paragraphs):
        for p in paragraphs:
            _convert_paragraph(p, user)

    fix_paragraphs(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                fix_paragraphs(cell.paragraphs)
    for section in doc.sections:
        for part in (section.header, section.footer):
            if part is not None:
                fix_paragraphs(part.paragraphs)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


@router.message(F.document)
async def on_document(message: Message, user: User) -> None:
    doc = message.document
    name = doc.file_name or "matn.txt"
    ext = Path(name).suffix.lower()
    if ext not in TEXT_EXT | DOCX_EXT:
        await message.answer(t("file_bad", user.ui))
        return
    if (doc.file_size or 0) > config.max_file_mb * 1024 * 1024:
        await message.answer(t("file_big", user.ui, mb=config.max_file_mb))
        return

    note = await message.answer(t("file_wait", user.ui))
    buf = io.BytesIO()
    await message.bot.download(doc, destination=buf)
    data = buf.getvalue()

    try:
        if ext in DOCX_EXT:
            payload = _convert_docx(data, user)
            chars = doc.file_size or len(data)
        else:
            text = _decode(data)
            payload = convert_for(user, text).text.encode("utf-8")
            chars = len(text)
    except ModuleNotFoundError:
        await note.edit_text(t("file_bad", user.ui))
        return

    new_name = f"{Path(name).stem}_{user.mode}{ext}"
    await message.answer_document(
        BufferedInputFile(payload, filename=new_name),
        caption=t("file_done", user.ui, chars=chars),
    )
    await note.delete()
    await db.bump(user.user_id, chars)
