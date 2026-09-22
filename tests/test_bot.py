"""Bot qatlami uçun oflayn testlar (Telegram serveri kerak emas)."""
from __future__ import annotations

import asyncio
import io
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from bot import db  # noqa: E402
from bot.handlers.documents import _convert_docx, _decode  # noqa: E402
from bot.keyboards import button_texts, main_kb, result_kb, settings_kb  # noqa: E402
from bot.texts import b, t  # noqa: E402
from bot.utils import as_html, convert_for, split_text  # noqa: E402

fails = 0


def check(cond, label, extra=""):
    global fails
    if not cond:
        fails += 1
    print(f"[{'OK ' if cond else 'XATO'}] {label}{(' — ' + str(extra)) if not cond else ''}")


async def main() -> None:
    tmp = Path(tempfile.mkdtemp()) / "test.sqlite3"
    await db.init(tmp)

    u = await db.get_user(777, "ali", "Ali Valiyev")
    check(u.mode == "yangi" and u.ui == "yangi", "yangi foydalanuvçi standart sozlamalari")

    await db.set_field(777, "mode", "kirill")
    u = await db.get_user(777)
    check(u.mode == "kirill", "sozlama saqlandi")

    await db.bump(777, 120)
    await db.bump(777, 30)
    u = await db.get_user(777)
    check(u.conv == 2 and u.chars == 150, "statistika ösdi", f"{u.conv}/{u.chars}")

    s = await db.stats()
    check(s["users"] == 1 and s["conv"] == 2 and s["today"] == 1, "umumiy statistika", s)

    await db.set_chat_auto(-100123, True, "yangi")
    auto, mode = await db.chat_auto(-100123)
    check(auto and mode == "yangi", "kanal avto-rejimi")

    ids = await db.all_user_ids()
    check(ids == [777], "foydalanuvçilar röyxati", ids)
    await db.mark_blocked(777)
    check(await db.all_user_ids() == [], "bloklangan foydalanuvçi çiqarildi")

    # --- konvertatsiya foydalanuvçi sozlamalari bilan ---
    u.mode, u.smart, u.keep_foreign, u.fmt = "yangi", 1, 1, "oddiy"
    check(convert_for(u, "O'zbekiston shahri").text == "Özbekiston şahri", "convert_for")
    u.keep_foreign = 0
    check(convert_for(u, "Chelsea").text == "Çelsea", "keep_foreign o'çirilgan")

    # --- uzun matnni bölaklaş ---
    long = ("Toshkent shahri juda go'zal. " * 400).strip()
    parts = split_text(long)
    check(all(len(p) <= 3900 for p in parts) and "".join(p for p in parts), "split_text çegara", len(parts))
    check(sum(len(p) for p in parts) >= len(long) - 2 * len(parts), "split_text matn yöqolmadi")

    # --- HTML xavfsizligi ---
    check(as_html("<b>ha & yo'q</b>") == "&lt;b&gt;ha &amp; yo'q&lt;/b&gt;", "HTML ekranlaş")
    check(as_html("salom", "kod") == "<code>salom</code>", "kod körinişi")

    # --- klaviaturalar ---
    kb = result_kb("ab1", "yangi", "oddiy", "yangi")
    labels = [btn.text for row in kb.inline_keyboard for btn in row]
    check(len(labels) == 4 and any("Eski" in x for x in labels), "natija klaviaturasi", labels)
    check(all(len(btn.callback_data.encode()) <= 64
              for row in kb.inline_keyboard for btn in row), "callback_data 64 baytdan oşmaydi")
    check(len(settings_kb(u).inline_keyboard) == 5, "sozlama klaviaturasi")
    check(len(main_kb("kirill").keyboard) == 3, "asosiy klaviatura")
    bt = button_texts()
    check("🇺🇿 Кирилл" in bt and bt["🇺🇿 Кирилл"] == "kirill", "kirill tugma matni tanildi")

    # --- interfeys matnlari uç yozuvda ham buzilmaydi ---
    for ui in ("yangi", "eski", "kirill"):
        txt = t("help", ui)
        check("/start" in txt and txt.count("<b>") == txt.count("</b>"),
              f"interfeys matni butun ({ui})")
    check("Kirill" in b("kirill", "yangi"), "tugma nomi")

    # --- kodlaşni aniqlaş ---
    cyr = "Ўзбекистон Республикаси Тошкент шахри барча учун"
    check(_decode(cyr.encode("cp1251")) == cyr, "cp1251 faylni öqiş", _decode(cyr.encode("cp1251")))
    check(_decode(cyr.encode("utf-8")) == cyr, "utf-8 kirill faylni öqiş")
    check(_decode(b"") == "", "böş fayl")
    check(_decode("Toshkent".encode("utf-8")) == "Toshkent", "utf-8 faylni öqiş")

    # --- docx ---
    try:
        import docx
        d = docx.Document()
        d.add_paragraph("O'zbekiston shahri")
        table = d.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Toshkent uchun"
        buf = io.BytesIO(); d.save(buf)
        u.mode, u.keep_foreign = "yangi", 1
        out = docx.Document(io.BytesIO(_convert_docx(buf.getvalue(), u)))
        got = out.paragraphs[0].text
        cell = out.tables[0].cell(0, 0).text
        check(got == "Özbekiston şahri", "docx paragraf", got)
        check(cell == "Toşkent uçun", "docx jadval", cell)
    except ModuleNotFoundError:
        print("[SKIP] docx testi — python-docx örnatilmagan")

    await db.close()
    print("\n" + ("BARÇASI OʻTDI" if not fails else f"{fails} ta test yiqildi"))
    raise SystemExit(1 if fails else 0)


asyncio.run(main())
